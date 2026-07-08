import os
from pathlib import Path


def extract_images_from_doc(file_path: str, output_dir: str) -> list[dict]:
    ext = Path(file_path).suffix.lower()
    os.makedirs(output_dir, exist_ok=True)

    if ext == ".pdf":
        return _extract_from_pdf(file_path, output_dir)
    elif ext == ".docx":
        return _extract_from_docx(file_path, output_dir)
    elif ext == ".pptx":
        return _extract_from_pptx(file_path, output_dir)
    else:
        return []


def _extract_from_pdf(file_path: str, output_dir: str) -> list[dict]:
    images = []
    try:
        import fitz
        doc = fitz.open(file_path)
        for page_num, page in enumerate(doc):
            for img_index, img in enumerate(page.get_images(full=True)):
                xref = img[0]
                base_image = doc.extract_image(xref)
                image_bytes = base_image["image"]
                image_ext = base_image["ext"]
                image_filename = f"pdf_page{page_num + 1}_img{img_index + 1}.{image_ext}"
                image_path = os.path.join(output_dir, image_filename)
                with open(image_path, "wb") as f:
                    f.write(image_bytes)
                images.append({
                    "filename": image_filename,
                    "path": image_path,
                    "page": page_num + 1,
                    "extension": image_ext,
                    "size": len(image_bytes),
                })
        doc.close()
    except ImportError:
        pass
    return images


def _extract_from_docx(file_path: str, output_dir: str) -> list[dict]:
    images = []
    try:
        from docx import Document
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        doc = Document(file_path)
        for i, rel in enumerate(doc.part.rels.values()):
            if "image" in rel.reltype:
                image = rel.target_part
                image_ext = Path(image.partname).suffix
                image_filename = f"docx_img{i + 1}{image_ext}"
                image_path = os.path.join(output_dir, image_filename)
                with open(image_path, "wb") as f:
                    f.write(image.blob)
                images.append({
                    "filename": image_filename,
                    "path": image_path,
                    "extension": image_ext,
                    "size": len(image.blob),
                })
    except ImportError:
        pass
    return images


def _extract_from_pptx(file_path: str, output_dir: str) -> list[dict]:
    images = []
    try:
        from pptx import Presentation
        prs = Presentation(file_path)
        img_idx = 0
        for slide_num, slide in enumerate(prs.slides):
            for shape in slide.shapes:
                if shape.shape_type == 13:
                    image = shape.image
                    image_ext = image.content_type.split("/")[-1]
                    if image_ext == "jpeg":
                        image_ext = "jpg"
                    image_filename = f"pptx_slide{slide_num + 1}_img{img_idx + 1}.{image_ext}"
                    image_path = os.path.join(output_dir, image_filename)
                    with open(image_path, "wb") as f:
                        f.write(image.blob)
                    images.append({
                        "filename": image_filename,
                        "path": image_path,
                        "slide": slide_num + 1,
                        "extension": image_ext,
                        "size": len(image.blob),
                    })
                    img_idx += 1
    except ImportError:
        pass
    return images
